import openpyxl,json,re
wb=openpyxl.load_workbook('sheet.xlsx',data_only=True);ws=wb['Modway Curation']
rows={}
for r in range(5,ws.max_row+1):
  v=[ws.cell(r,c).value for c in range(1,17)]
  if not any(v[2:4]) and not v[13]: continue
  if (v[0] or '').strip()=='Live': continue
  rows[r]={'A':v[0],'C':v[2],'D':v[3],'H':v[7],'J':v[9],'K':v[10],'L':v[11],'N':(v[13] or '').strip(),'O':v[14],'P':v[15]}
json.dump(rows,open('rows.json','w'),ensure_ascii=False)
skus=[x['N'] for x in rows.values()]
print(len(rows),'rows',len(set(skus)),'skus; blank sku',sum(1 for s in skus if not s))
bases=sorted({re.match(r'([A-Z]+-\d+)',s).group(1) for s in skus if re.match(r'([A-Z]+-\d+)',s)})
print(len(bases),'bases')
json.dump(bases,open('bases.json','w'))
